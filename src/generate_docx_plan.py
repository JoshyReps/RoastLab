import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = docx.Document()

# Styles
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(11)
normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

title = doc.add_paragraph()
title_run = title.add_run('RoastLab Coffee Co. — Machine Learning Finals Project Plan')
title_run.bold = True
title_run.font.size = Pt(20)
title_run.font.color.rgb = RGBColor(0x3E, 0x27, 0x23)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

sub = doc.add_paragraph()
sub_run = sub.add_run('ELEC026.23 Machine Learning | Completed Specification & Prototype Documentation')
sub_run.italic = True
sub_run.font.size = Pt(12)
sub_run.font.color.rgb = RGBColor(0x79, 0x55, 0x48)
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Group Members: Reponoya, Olpenado, Palencia, Rivera\nDate: October 2026\n')

# Company Profile
doc.add_heading('Company Profile: RoastLab Coffee Co.', level=1)
doc.add_paragraph('Industry: Specialty Food & Beverage / Smart Retail')
doc.add_paragraph(
    'Company Description: RoastLab Coffee Co. is a modern artisanal cafe chain that '
    'integrates practical machine learning tools into daily store operations to streamline '
    'inventory preparation, personalize customer beverage choices, and monitor customer '
    'feedback in real time.'
)

doc.add_paragraph(
    'Vision: To deliver an unmatched artisanal coffee experience by seamlessly pairing '
    'hand-crafted brewing with intelligent, data-driven cafe management.'
)
doc.add_paragraph(
    'Mission: To empower store managers and baristas with transparent machine learning '
    'tools that reduce food waste, simplify drink selection for guests, and ensure consistent '
    'service quality across all branches.'
)

p_goals = doc.add_paragraph()
p_goals_run = p_goals.add_run('Operational Goals:')
p_goals_run.bold = True
doc.add_paragraph('• Accurately forecast daily batch production to minimize ingredient waste and prevent stockouts.', style='List Bullet')
doc.add_paragraph('• Deliver personalized drink and coffee bean recommendations to increase customer satisfaction and boost sales.', style='List Bullet')
doc.add_paragraph('• Automate review sentiment analysis to identify and resolve customer service issues in real time.', style='List Bullet')
doc.add_paragraph('• Ensure precise roasting quality control across batches using deep sensory profile prediction.', style='List Bullet')

# 1. Linear Regression
doc.add_heading('1. Linear Regression Plan', level=2)
doc.add_paragraph(
    "Idea: Uses ordinary least squares linear regression (Y = mx + b) to predict tomorrow's cold brew liters (Y) "
    "based on forecasted peak temperature (X). It tells baristas the exact batch size to steep 24 hours in "
    "advance—preventing stockouts on hot days and coffee waste on cold ones."
)
doc.add_paragraph('Daily Temperature to Cold Brew Demand:')
doc.add_paragraph('• X (Input): Forecasted daily high temperature (°C).', style='List Bullet')
doc.add_paragraph('• Y (Target): Total liters of cold brew needed.', style='List Bullet')
doc.add_paragraph('Business Value: Determines cold brew batch sizes 24 hours in advance to prevent running out on hot days or wasting unsold stock on cold days.')
doc.add_paragraph(
    'App Feature: Users adjust a temperature slider for tomorrow\'s weather forecast, and the system '
    'instantly outputs the exact liters of cold brew to steep (e.g., "30°C -> Steep 24.3 Liters"), '
    'complete with an interactive scatter plot and fitted regression line.'
)

# 2. Naïve Bayes
doc.add_heading('2. Naïve Bayes Classifier Plan', level=2)
doc.add_paragraph(
    'Idea: Uses Multinomial Naïve Bayes classification to predict whether a customer prefers a Hot or Iced '
    'beverage based on contextual order conditions. It multiplies conditional probability ratios from past '
    'orders and applies Laplace smoothing (+1) so new or rare combinations do not break the calculation.'
)
doc.add_paragraph('Categorical Inputs:')
doc.add_paragraph('• Time of Day (Morning, Afternoon, Evening)', style='List Bullet')
doc.add_paragraph('• Order Type (Dine-in, Takeaway)', style='List Bullet')
doc.add_paragraph('• Outdoor Weather (Warm, Cool)', style='List Bullet')
doc.add_paragraph('• Predicted Class (Y): Drink Preference (Hot vs. Iced)', style='List Bullet')
doc.add_paragraph('Business Value: Speeds up order preparation and personalizes app menu recommendations to boost customer satisfaction and drive upsells.')
doc.add_paragraph(
    'App Feature: Users select conditions from three simple dropdowns, and the system instantly outputs '
    'the recommended drink type with a confidence score and probability bar chart (e.g., "78% Probability of Iced Beverage").'
)

# 3. KNN
doc.add_heading('3. K-Nearest Neighbor (KNN) Plan', level=2)
doc.add_paragraph(
    'Idea: Uses K-Nearest Neighbors (k = 3) to recommend a coffee drink by calculating the Euclidean distance '
    'd(P, Q) = sqrt((x2 - x1)^2 + (y2 - y1)^2) between a customer\'s taste preferences and existing menu items.'
)
doc.add_paragraph('Continuous Inputs:')
doc.add_paragraph('• Sweetness Level (Scale 1 to 10)', style='List Bullet')
doc.add_paragraph('• Caffeine Strength (Scale 1 to 10)', style='List Bullet')
doc.add_paragraph('Predicted Class (Y): Recommended Menu Drink (e.g., Signature Cold Brew, Caramel Macchiato, Flat White, Artisanal Americano).')
doc.add_paragraph('Business Value: Enhances product discovery and increases sales and user satisfaction by delivering personalized drink recommendations.')
doc.add_paragraph(
    'App Feature: Users adjust two taste sliders (Sweetness & Caffeine) on an interactive 2D chart, '
    'and the app instantly highlights the 3 nearest matching drinks on the menu with exact distance metrics.'
)

# 4. SVM
doc.add_heading('4. Support Vector Machine (SVM) Plan', level=2)
doc.add_paragraph(
    'Idea: Uses a Support Vector Machine (SVM) with a linear kernel to classify customer review text into '
    'sentiment categories by establishing an optimal maximum-margin decision boundary between positive and negative '
    'word features extracted via TF-IDF.'
)
doc.add_paragraph('Text Inputs:')
doc.add_paragraph('• Customer Review Text (e.g., "The cold brew was sour and the service was slow")', style='List Bullet')
doc.add_paragraph('Predicted Class (Y): Review Sentiment (Positive vs. Negative).')
doc.add_paragraph('Business Value: Automatically alerts branch managers about negative feedback in real time to resolve service issues quickly and protect customer loyalty.')
doc.add_paragraph(
    'App Feature: Users type a quick review into a text box, and the app instantly flags it as Positive (Green) '
    'or Negative (Red) with margin confidence score to route bad reviews directly to customer support.'
)

# 5. Decision Tree
doc.add_heading('5. Decision Tree Plan', level=2)
doc.add_paragraph(
    'Idea: Uses an ID3 Decision Tree algorithm to match customers with the perfect Single-Origin Coffee Bean '
    'by calculating Entropy and Information Gain across their taste profile preferences.'
)
doc.add_paragraph('Categorical Inputs:')
doc.add_paragraph('• Body Preference (Light/Tea-like, Medium/Smooth, Heavy/Syrupy)', style='List Bullet')
doc.add_paragraph('• Acidity Preference (High/Bright, Balanced, Low/Mild)', style='List Bullet')
doc.add_paragraph('• Primary Flavor Profile (Fruity/Floral, Chocolatey/Nutty, Earthy/Spicy)', style='List Bullet')
doc.add_paragraph('Predicted Class (Y): Recommended Bean Origin (e.g., Ethiopian Yirgacheffe, Kenyan AA, Colombian Supremo, Sumatra Mandheling).')
doc.add_paragraph('Business Value: Drives sales of high-margin specialty coffee beans and helps everyday customers confidently order premium pour-overs without feeling intimidated by coffee jargon.')
doc.add_paragraph(
    'App Feature: Customers take a 3-question "Bean Selector Quiz" on the app or ordering kiosk, '
    'and the decision tree branches immediately output their ideal single-origin pour-over match along with elevation, region, and tasting notes.'
)

# 6. ANN
doc.add_heading('6. Artificial Neural Network (ANN) Plan [COMPLETED]', level=2)
doc.add_paragraph(
    'Idea: Uses a Multi-Layer Perceptron (MLP) Artificial Neural Network with ReLU hidden layers to model non-linear roasting '
    'thermodynamics. It predicts the final Specialty Coffee Association (SCA) Cupping Score (80 to 95 scale) and batch quality '
    'tier directly from drum sensor telemetry before the batch finishes cooling.'
)
doc.add_paragraph('Continuous Sensor Inputs (5 Features):')
doc.add_paragraph('• Charge Temperature (°C): Initial drum temperature when green beans are dropped (190°C - 215°C)', style='List Bullet')
doc.add_paragraph('• Drop Temperature (°C): Final roast temperature at bean discharge (205°C - 228°C)', style='List Bullet')
doc.add_paragraph('• Total Roast Duration (min): Overall duration inside the drum roaster (9.0 - 14.5 min)', style='List Bullet')
doc.add_paragraph('• First Crack Time (min): Exothermic reaction milestone when bean cell walls expand (7.0 - 11.5 min)', style='List Bullet')
doc.add_paragraph('• Average Rate of Rise (RoR in °C/min): Speed of temperature increase during caramelization (7.0 - 14.0 °C/min)', style='List Bullet')
doc.add_paragraph('Target (Y): Predicted SCA Cupping Score (80 - 95 scale) and Quality Tier (Presidential Specialty, Specialty Grade, Premium Commercial, Underdeveloped).')
doc.add_paragraph('Network Architecture: Input [5] -> Hidden Layer 1 [32 neurons, ReLU] -> Hidden Layer 2 [16 neurons, ReLU] -> Output [1 linear neuron, Cupping Score]. Trained using the Adam optimizer.')
doc.add_paragraph('Business Value: Prevents defective or underdeveloped coffee batches from reaching store shelves, saves green bean costs, and guarantees brand consistency across all franchise branches.')
doc.add_paragraph(
    'App Feature: Roasters and managers adjust roaster telemetry sliders, and the neural network instantly outputs '
    'the predicted cupping score, quality badge, and animated score gauge.'
)

# Summary Table
doc.add_heading('Algorithm Benchmark & Metrics Summary', level=2)
table = doc.add_table(rows=1, cols=5)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Algorithm'
hdr_cells[1].text = 'Business Function'
hdr_cells[2].text = 'Key Metric'
hdr_cells[3].text = 'Error / Criteria'
hdr_cells[4].text = 'Dataset Size'

rows_data = [
    ('Linear Regression', 'Cold Brew Volume Demand', 'R² = 0.909', 'RMSE = 1.78 L', '180 days'),
    ('Naïve Bayes', 'Hot vs. Iced Prediction', 'Accuracy = 81.6%', 'Laplace α = 1.0', '320 orders'),
    ('KNN (k=3)', 'Drink Menu Recommender', 'k = 3 Neighbors', 'Euclidean Metric', '350 samples'),
    ('Linear SVM', 'Review Sentiment Analysis', 'Accuracy = 100%', '24 Support Vectors', '24 reviews'),
    ('Decision Tree (ID3)', 'Single-Origin Bean Match', 'Tree Depth = 3', 'Entropy / Gain', '15 profiles'),
    ('ANN (MLP)', 'Roast Cupping QC Score', 'R² = 0.987', 'RMSE = 0.38 pts', '350 roasts')
]

for row in rows_data:
    row_cells = table.add_row().cells
    for i, val in enumerate(row):
        row_cells[i].text = val

out_path = r'docs\Final_Machine-Learning-Plan_Completed.docx'
doc.save(out_path)
print(f'Successfully generated {out_path}')
